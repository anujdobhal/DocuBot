from pathlib import Path
import fitz
from docx import Document

def extract_text(path: Path) -> str:
    ext = path.suffix.lower()

    if ext == ".pdf":
        with fitz.open(path) as pdf:
            return "\n\n".join(page.get_text("text") for page in pdf)

    if ext == ".docx":
        doc = Document(path)
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text.strip() for cell in row.cells))
        return "\n".join(parts)

    if ext == ".txt":
        return path.read_text(encoding="utf-8", errors="replace")

    raise ValueError(f"Unsupported file type: {ext}")
